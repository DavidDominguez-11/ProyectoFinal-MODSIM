"""Inspecciona la hoja de observaciones del Camp Fire publicada por NIST.

No modifica ni corrige el archivo fuente. En particular, informa los años
que aparecen en su columna Date para detectar inconsistencias con 2018.
"""

from __future__ import annotations

import collections
import hashlib
import io
import json
import urllib.request

import openpyxl


URL = (
    "https://data.nist.gov/od/ds/mds2-2954/"
    "NIST%20TN%202135%20Appendix%20F%20Fire%20Spread%20Data.xlsx"
)


def main() -> None:
    with urllib.request.urlopen(URL, timeout=60) as response:
        contents = response.read()
    workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)

    year_counts: collections.Counter[str] = collections.Counter()
    fire_types: collections.Counter[str] = collections.Counter()
    sheet_counts: dict[str, int] = {}
    source_ids: set[str] = set()
    unique_observations: set[tuple[object, ...]] = set()
    valid_coordinates = 0
    total_rows = 0
    malformed_rows = 0
    time_windows: list[float] = []

    for sheet in workbook:
        nonempty = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue
            nonempty += 1
            total_rows += 1
            source_id, date, time, _, window, latitude, longitude, *_ = row
            if hasattr(date, "year"):
                year_counts[str(date.year)] += 1
            elif date is not None:
                year_counts[str(date)] += 1
            if source_id:
                source_ids.add(str(source_id))
            fire_type = row[10] if len(row) > 10 else None
            if fire_type:
                fire_types[str(fire_type)] += 1
            if isinstance(latitude, (int, float)) and isinstance(longitude, (int, float)):
                valid_coordinates += 1
            else:
                malformed_rows += 1
            if isinstance(window, (int, float)):
                time_windows.append(float(window))
            unique_observations.add(
                (str(source_id), str(date), str(time), latitude, longitude, str(fire_type))
            )
        sheet_counts[sheet.title] = nonempty

    result = {
        "url": URL,
        "download_bytes": len(contents),
        "sha256": hashlib.sha256(contents).hexdigest(),
        "sheets": len(workbook.sheetnames),
        "nonempty_rows_by_sheet": sheet_counts,
        "row_occurrences_across_sheets": total_rows,
        "unique_source_ids": len(source_ids),
        "unique_observation_keys": len(unique_observations),
        "year_values": dict(year_counts),
        "type_codes": dict(fire_types),
        "rows_with_numeric_coordinates": valid_coordinates,
        "rows_without_numeric_coordinates": malformed_rows,
        "time_window_minutes": {
            "count": len(time_windows),
            "min": min(time_windows) if time_windows else None,
            "max": max(time_windows) if time_windows else None,
        },
        "warning": (
            "Los años se muestran tal como aparecen en el libro; "
            "no se debe corregir Date ni contar hojas repetidas como observaciones nuevas."
        ),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
